from pathlib import Path
from itertools import product
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
root=Path(r'C:\laragon\www\SUSTAINOVATA-main')
out=root/'output/doc/Aktivitas_Kelas_Strong_Robust_Rusdany_202210370311406.docx'
out.parent.mkdir(parents=True,exist_ok=True)
d=Document();s=d.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69)
s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.7)
for sn in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[sn];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Title'].font.size=Pt(21)
d.styles['Heading 1'].font.size=Pt(14)
d.styles['Heading 2'].font.size=Pt(11)
for st in d.styles:
 for e in list(st.element.findall('.//' + qn('w:pBdr'))):e.getparent().remove(e)
def p(x):d.add_paragraph(x)
def h(x):d.add_heading(x,1)
def table(headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=x
 for vals in rows:
  for c,x in zip(t.add_row().cells,vals):c.text=str(x)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[j])
   tcp=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tcp.append(b)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF5');tcp.append(sh)
   for pa in c.paragraphs:
    pa.paragraph_format.space_after=Pt(4);pa.paragraph_format.space_before=Pt(4)
    for r in pa.runs:r.font.size=Pt(9);r.bold=i==0
 return t
d.add_paragraph('Pengujian Kelas Ekuivalensi Strong Robust',style='Title')
p('Sustainova Web | Fitur Register | Validasi Nama Lengkap dan Alamat')
p('Rusdany Maestro | 202210370311406 | Kelas C | 15 Oktober 2026')
h('1 Tujuan dan ruang lingkup')
p('Menyusun test case Strong Robust untuk dua input yang diketik, yaitu Nama Lengkap dan Alamat. Seluruh kombinasi kelas valid dan tidak valid kedua input dicakup. Kolom lainnya dipertahankan valid sebagai kondisi kontrol. Desain ini menguji validasi frontend pada /register, bukan seluruh fitur Register atau keseluruhan aplikasi.')
h('2 Dasar aturan input')
p('Aturan diperiksa langsung pada frontend/src/page/auth/register.vue: validateField untuk name mulai baris 169 dan address mulai baris 201. Kedua kolom memakai required dan validasi saat blur. handleRegister berhenti jika validateForm gagal; tombol Daftar dinonaktifkan saat hasErrors bernilai true.')
table(['ID','Aturan yang menjadi dasar pengujian'],[
 ('RQ-N','Nama wajib diisi dan panjangnya minimal 3 karakter.'),
 ('RQ-A','Alamat wajib diisi dan panjangnya minimal 10 karakter.'),
 ('RQ-F','Jika salah satu input tidak valid, form tidak boleh mengirim pendaftaran; jika keduanya valid dan kolom kontrol valid, validasi mengizinkan pengiriman.')],[.65,6.2])
h('3 Kelas ekuivalensi')
table(['Input','Kelas','Kondisi panjang','Kategori','Perwakilan'],[
 ('Nama','N-V','n >= 3','Valid','Rusdany (7)'),('Nama','N-I1','n = 0','Tidak valid','Kosong (0)'),('Nama','N-I2','1 <= n < 3','Tidak valid','Ru (2)'),
 ('Alamat','A-V','a >= 10','Valid','Jalan Melati 10 (15)'),('Alamat','A-I1','a = 0','Tidak valid','Kosong (0)'),('Alamat','A-I2','1 <= a < 10','Tidak valid','Malang (6)')],[.65,.7,1.3,1,3.2])
p('n dan a adalah panjang string menurut JavaScript. Contoh memakai karakter ASCII. Label valid berarti lolos aturan wajib/panjang pada UI. Batas maksimum, aturan hanya huruf, whitespace-only, karakter khusus, dan validasi backend tidak diklaim tercakup karena tidak menjadi aturan yang dimodelkan dalam latihan ini.')
d.add_page_break()
h('4 Perhitungan dan matriks Strong Robust')
p('Nama memiliki 1 kelas valid + 2 kelas tidak valid. Alamat juga memiliki 1 kelas valid + 2 kelas tidak valid. Jumlah test case = (1 + 2) x (1 + 2) = 9. Strong berarti seluruh kombinasi antarkelas; Robust berarti memasukkan kelas tidak valid, termasuk beberapa input tidak valid sekaligus.')
names=[('N-V','Rusdany'),('N-I1',''),('N-I2','Ru')]
addresses=[('A-V','Jalan Melati 10'),('A-I1',''),('A-I2','Malang')]
rows=[]
for i,((nc,nv),(ac,av)) in enumerate(product(names,addresses),1):
 exp='Validasi mengizinkan pengiriman.' if nc=='N-V' and ac=='A-V' else 'Pendaftaran diblokir; '
 if exp.startswith('Pendaftaran'):
  errs=[]
  if nc=='N-I1':errs.append('nama kosong')
  if nc=='N-I2':errs.append('nama < 3')
  if ac=='A-I1':errs.append('alamat kosong')
  if ac=='A-I2':errs.append('alamat < 10')
  exp+=', '.join(errs)+'.'
 rows.append((f'SR-{i:02}',nc+' / '+ac,nv or '(kosong)',av or '(kosong)',exp))
table(['ID','Kombinasi','Nama','Alamat','Hasil yang diharapkan'],rows,[.55,1.15,.8,1.4,2.95])
p('Semua test case berstatus Belum diuji. Tabel ini adalah desain pengujian; hasil aktual dan status Lolos/Gagal baru diisi setelah eksekusi. Penolakan input tidak valid yang sesuai harapan berarti test case Lolos, bukan Gagal.')
h('5 Kondisi kontrol dan keputusan hasil')
p('Setiap kasus dimulai dari form bersih. Kolom kontrol: email unik yang belum terdaftar, misalnya sr01-20261015@example.test; telepon 081234567890; password Latihan123!; konfirmasi password Latihan123!. Semua data merupakan data dummy. Jika kasus valid diulang, gunakan email baru agar duplikasi email tidak memengaruhi hasil.')
p('SR-01: tidak ada error nama/alamat dan form boleh mengirim permintaan. Keberhasilan pembuatan akun selanjutnya bergantung pada backend/database serta berada di luar oracle validasi frontend latihan ini. SR-02 sampai SR-09: input ditolak dan tidak ada pengiriman pendaftaran karena validasi.')
d.add_page_break()
h('6 Langkah eksekusi manual')
for text in [
 '1. Buka http://localhost:5173/register. Muat ulang sebelum memulai setiap kasus agar error dan input kasus sebelumnya tidak terbawa.',
 '2. Isi kolom kontrol dengan data valid. Gunakan email unik untuk percobaan yang dapat mengirim permintaan.',
 '3. Ketik Nama sesuai baris SR yang sedang diuji. Tekan Tab atau klik kolom lain untuk memicu blur. Untuk kelas kosong, fokuskan kolom, pastikan kosong, lalu keluar dari kolom.',
 '4. Ketik Alamat sesuai baris yang sama. Keluar dari kolom untuk memicu validasi. Periksa kedua kolom, khususnya pada SR-05, SR-06, SR-08, dan SR-09 yang memiliki dua input tidak valid.',
 '5. Amati pesan pada setiap kolom dan apakah tombol Daftar dinonaktifkan. Jangan memaksa klik tombol disabled. Browser juga dapat menolak kolom required yang kosong sebelum handler pengiriman dijalankan.',
 '6. Untuk SR-01, pastikan seluruh kolom valid. Jika eksekusi pengiriman dilakukan, klik Daftar dan gunakan panel Network browser untuk memeriksa permintaan POST /api/clients/register. Untuk kasus tidak valid, pastikan tidak ada permintaan baru dari percobaan tersebut.',
 '7. Simpan screenshot input dan pesan, tulis hasil aktual, lalu tentukan Lolos atau Gagal dengan membandingkannya terhadap harapan. Jika server tidak tersedia, pisahkan hasil validasi UI dari hasil pengiriman yang terblokir lingkungan.'
]:p(text)
h('7 Pesan validasi yang diperiksa')
table(['Kelas tidak valid','Pesan setelah blur'],[
 ('N-I1','Nama tidak boleh kosong'),('N-I2','Nama minimal 3 karakter'),('A-I1','Alamat tidak boleh kosong'),('A-I2','Alamat terlalu pendek, minimal 10 karakter')],[1.3,5.55])
h('8 Pemeriksaan kelengkapan dan penjelasan kelas')
p('Cakupan kombinasi desain: 9 dari 9 kombinasi. Terdapat 1 kombinasi seluruh input valid, 4 kombinasi satu input tidak valid, dan 4 kombinasi dua input tidak valid. Empat kasus terakhir membedakan desain ini dari Weak Robust. Angka cakupan desain bukan hasil eksekusi pengujian.')
p('Slide kelas juga menyebut Weak Normal. Karena masing-masing input mempunyai satu kelas valid dalam model ini, Weak Normal memerlukan satu kombinasi, yaitu N-V x A-V. Kombinasi tersebut sudah diwakili SR-01; tidak diperlukan test case tambahan atau pengujian ganda. Teknik utama yang dipilih tetap Strong Robust.')
p('Penjelasan saat presentasi: Saya memilih input Nama dan Alamat pada Register. Nama minimal 3 dan alamat minimal 10 karakter. Masing-masing saya bagi menjadi valid, kosong, dan terlalu pendek. Seluruh kombinasinya berjumlah sembilan, termasuk dua input tidak valid bersamaan. Hasil aktual belum saya klaim sebelum eksekusi.')
p('Rujukan konsep: materi Equivalence Class Testing, University of Kansas, bagian 3.4 Strong Robust Equivalence Class Testing: https://people.eecs.ku.edu/~saiedian/Teaching/814/Readings/Davies-EC-Testing-Notes.pdf. Dasar aturan aplikasi: frontend/src/page/auth/register.vue, validateField dan handleRegister, diperiksa 15 Oktober 2026.')
assert len(rows)==9 and len({r[1] for r in rows})==9
assert sum(n=='N-V' and a=='A-V' for (n,_),(a,__) in product(names,addresses))==1
assert len('Jalan Melati 10')==15
d.save(out)
print(out)
