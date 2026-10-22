from docx import Document
from lxml import etree as E
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
d=Document(r'output/doc/PKPL_Modul_1_Test_Plan_Admin_Rusdany_202210370311406_Revisi_v1_2.docx')
for i,p in enumerate(d.paragraphs):
 if 20<=i<=35 or 58<=i<=68:
  pr=p._p.find('w:pPr',ns)
  print(i,repr(p.text[:50]),'pPr',[(E.QName(e).localname,dict(e.attrib)) for e in pr] if pr is not None else [],'breaks',[dict(e.attrib) for e in p._p.findall('.//w:br',ns)])
t=Document(r'output/doc/PKPL_Modul_2_LAB_TDS_TTM_Admin_Rusdany_202210370311406.docx')
for table in t.tables:
 print('HEADER', [E.tostring(e).decode()[-200:] for e in table.rows[0]._tr.findall('.//w:tblHeader',ns)])
